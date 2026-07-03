"""AI Agent Studio (Phase 1): create/train/test intelligent AI agents.
Each agent has its own prompts + private knowledge base (RAG via TF-IDF retrieval in MongoDB).
Powers a real, knowledge-grounded conversation engine (chat playground + AI call generation).
"""
import os
import io
import re
import math
import uuid
import json
import logging
from datetime import datetime, timezone
from collections import Counter

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from pydantic import BaseModel, Field
from typing import Optional, List

logger = logging.getLogger(__name__)

PERSONALITIES = ["professional", "friendly", "formal", "sales", "consultative",
                 "empathetic", "supportive", "technical", "confident", "motivational"]
CATEGORIES = ["Sales Agent", "Customer Support", "Lead Qualification", "Appointment Booking",
              "Admission Counsellor", "Follow-up Agent", "Retention Agent", "Cold Calling", "Custom"]


# ---------------- Models ----------------
class AgentIn(BaseModel):
    name: str
    description: Optional[str] = ""
    category: str = "Custom"
    industry: Optional[str] = ""
    language: str = "english"
    personality: str = "professional"
    temperature: float = 0.6
    goal: Optional[str] = ""
    greeting: Optional[str] = ""
    closing: Optional[str] = ""
    fallback: Optional[str] = "I'm not sure about that — let me connect you with a human colleague."
    system_prompt: Optional[str] = ""
    voice: Optional[str] = "female"
    status: str = "draft"


class ChatIn(BaseModel):
    message: str
    session_id: Optional[str] = None


class TextKnowledgeIn(BaseModel):
    title: str
    content: str


# ---------------- Text extraction ----------------
def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            return "\n".join((p.extract_text() or "") for p in reader.pages)
        if name.endswith(".docx"):
            from docx import Document
            doc = Document(io.BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs)
        if name.endswith((".xlsx", ".xls")):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    out.append(" ".join(str(c) for c in row if c is not None))
            return "\n".join(out)
        # txt / csv / md / html / json
        return data.decode("utf-8", errors="ignore")
    except Exception as e:
        logger.error(f"extract_text failed for {filename}: {e}")
        return data.decode("utf-8", errors="ignore") if data else ""


def chunk_text(text: str, size: int = 900, overlap: int = 150) -> List[str]:
    text = re.sub(r"\s+", " ", text or "").strip()
    if not text:
        return []
    chunks, i = [], 0
    while i < len(text):
        chunks.append(text[i:i + size])
        i += size - overlap
    return chunks


# ---------------- TF-IDF retrieval ----------------
_WORD = re.compile(r"[a-z0-9]+")


def _tokens(s: str) -> List[str]:
    return _WORD.findall((s or "").lower())


async def retrieve(db, agent_id: str, query: str, k: int = 5) -> List[dict]:
    docs = await db.agent_knowledge.find({"agent_id": agent_id}).to_list(2000)
    if not docs:
        return []
    q_terms = set(_tokens(query))
    if not q_terms:
        return docs[:k]
    N = len(docs)
    df = Counter()
    doc_tokens = []
    for d in docs:
        toks = _tokens(d["content"])
        doc_tokens.append(toks)
        for t in set(toks):
            df[t] += 1
    scored = []
    for d, toks in zip(docs, doc_tokens):
        if not toks:
            continue
        tf = Counter(toks)
        score = 0.0
        for t in q_terms:
            if t in tf:
                idf = math.log((N + 1) / (df[t] + 1)) + 1
                score += (tf[t] / len(toks)) * idf
        if score > 0:
            scored.append((score, d))
    scored.sort(key=lambda x: x[0], reverse=True)
    top = scored[:k] if scored else [(0, d) for d in docs[:k]]
    maxs = top[0][0] or 1
    return [{"content": d["content"], "title": d.get("title", "doc"),
             "score": round(s / maxs, 3)} for s, d in top]


# ---------------- LLM ----------------
def _llm(system_message: str, session_id: str):
    from emergentintegrations.llm.chat import LlmChat
    key = os.environ["EMERGENT_LLM_KEY"]
    return LlmChat(api_key=key, session_id=session_id, system_message=system_message).with_model("openai", "gpt-4.1-mini")


def _agent_system_prompt(agent: dict, context: str) -> str:
    base = agent.get("system_prompt") or ""
    parts = [
        f"You are {agent['name']}, a {agent.get('personality','professional')} {agent.get('category','AI')} for "
        f"{agent.get('industry') or 'a business'}.",
        f"Your goal: {agent.get('goal') or 'help the customer and move them forward'}.",
        f"Always respond in {agent.get('language','english')}.",
        base,
        "Use ONLY the knowledge below to answer factual questions. If the answer isn't in the knowledge, "
        f"say: \"{agent.get('fallback')}\" — never invent facts, prices or policies.",
        "\n=== KNOWLEDGE ===\n" + (context or "(no knowledge provided yet)"),
    ]
    return "\n".join(p for p in parts if p)


async def _record_unknown(db, agent_id: str, question: str):
    q = (question or "").strip()
    if not q:
        return
    await db.agent_unknowns.update_one(
        {"agent_id": agent_id, "question": q.lower()},
        {"$setOnInsert": {"agent_id": agent_id, "question": q.lower(), "question_text": q,
                          "created_at": datetime.now(timezone.utc).isoformat(), "resolved": False},
         "$inc": {"count": 1}},
        upsert=True)


# ---------------- Router ----------------
def build_agents_router(db, get_current_user):
    router = APIRouter(prefix="/api/agents", tags=["agents"])

    def out(a: dict) -> dict:
        a.pop("_id", None)
        return a

    @router.get("")
    async def list_agents(user=Depends(get_current_user)):
        q = {} if user.get("role") == "admin" else {"owner_id": user["id"]}
        docs = await db.ai_agents.find(q).sort("created_at", -1).to_list(500)
        for d in docs:
            d.pop("_id", None)
            d["knowledge_count"] = await db.agent_knowledge.count_documents({"agent_id": d["id"]})
        return docs

    @router.post("")
    async def create_agent(payload: AgentIn, user=Depends(get_current_user)):
        now = datetime.now(timezone.utc).isoformat()
        doc = {"id": str(uuid.uuid4()), **payload.model_dump(),
               "owner_id": user["id"], "owner_name": user["name"],
               "created_at": now, "updated_at": now}
        await db.ai_agents.insert_one(dict(doc))
        return out(doc)

    @router.get("/meta")
    async def meta(user=Depends(get_current_user)):
        return {"personalities": PERSONALITIES, "categories": CATEGORIES}

    @router.get("/{agent_id}")
    async def get_agent(agent_id: str, user=Depends(get_current_user)):
        a = await db.ai_agents.find_one({"id": agent_id})
        if not a:
            raise HTTPException(404, "Agent not found")
        a = out(a)
        a["knowledge_count"] = await db.agent_knowledge.count_documents({"agent_id": agent_id})
        return a

    @router.put("/{agent_id}")
    async def update_agent(agent_id: str, payload: AgentIn, user=Depends(get_current_user)):
        existing = await db.ai_agents.find_one({"id": agent_id})
        if not existing:
            raise HTTPException(404, "Agent not found")
        new_prompt = payload.system_prompt or ""
        if (existing.get("system_prompt") or "") != new_prompt:
            await db.agent_prompt_versions.insert_one({
                "id": str(uuid.uuid4()), "agent_id": agent_id,
                "system_prompt": existing.get("system_prompt") or "",
                "created_at": datetime.now(timezone.utc).isoformat(),
                "created_by": user.get("name")})
        upd = {**payload.model_dump(), "updated_at": datetime.now(timezone.utc).isoformat()}
        r = await db.ai_agents.find_one_and_update({"id": agent_id}, {"$set": upd}, return_document=True)
        return out(r)

    @router.delete("/{agent_id}")
    async def delete_agent(agent_id: str, user=Depends(get_current_user)):
        await db.ai_agents.delete_one({"id": agent_id})
        await db.agent_knowledge.delete_many({"agent_id": agent_id})
        await db.agent_docs.delete_many({"agent_id": agent_id})
        return {"ok": True}

    # ---- Knowledge base ----
    @router.get("/{agent_id}/knowledge")
    async def list_knowledge(agent_id: str, user=Depends(get_current_user)):
        docs = await db.agent_docs.find({"agent_id": agent_id}).sort("created_at", -1).to_list(500)
        for d in docs:
            d.pop("_id", None)
        return docs

    async def _store_source(agent_id, title, text, source_type):
        chunks = chunk_text(text)
        doc_id = str(uuid.uuid4())
        now = datetime.now(timezone.utc).isoformat()
        for c in chunks:
            await db.agent_knowledge.insert_one({
                "id": str(uuid.uuid4()), "agent_id": agent_id, "doc_id": doc_id,
                "title": title, "content": c, "created_at": now})
        await db.agent_docs.insert_one({
            "id": doc_id, "agent_id": agent_id, "title": title, "type": source_type,
            "chunks": len(chunks), "chars": len(text), "created_at": now})
        return {"doc_id": doc_id, "chunks": len(chunks), "chars": len(text)}

    @router.post("/{agent_id}/knowledge/upload")
    async def upload_knowledge(agent_id: str, file: UploadFile = File(...), user=Depends(get_current_user)):
        if not await db.ai_agents.find_one({"id": agent_id}):
            raise HTTPException(404, "Agent not found")
        data = await file.read()
        text = extract_text(file.filename, data)
        if not text.strip():
            raise HTTPException(400, "Could not extract any text from this file")
        res = await _store_source(agent_id, file.filename, text, "file")
        return {"ok": True, **res}

    @router.post("/{agent_id}/knowledge/text")
    async def add_text_knowledge(agent_id: str, payload: TextKnowledgeIn, user=Depends(get_current_user)):
        if not await db.ai_agents.find_one({"id": agent_id}):
            raise HTTPException(404, "Agent not found")
        res = await _store_source(agent_id, payload.title, payload.content, "text")
        return {"ok": True, **res}

    @router.delete("/{agent_id}/knowledge/{doc_id}")
    async def delete_knowledge(agent_id: str, doc_id: str, user=Depends(get_current_user)):
        await db.agent_knowledge.delete_many({"agent_id": agent_id, "doc_id": doc_id})
        await db.agent_docs.delete_one({"id": doc_id, "agent_id": agent_id})
        return {"ok": True}

    # ---- Playground chat ----
    @router.post("/{agent_id}/chat")
    async def chat(agent_id: str, payload: ChatIn, user=Depends(get_current_user)):
        agent = await db.ai_agents.find_one({"id": agent_id})
        if not agent:
            raise HTTPException(404, "Agent not found")
        sources = await retrieve(db, agent_id, payload.message, k=5)
        context = "\n---\n".join(f"[{s['title']}] {s['content']}" for s in sources)
        sys = _agent_system_prompt(agent, context)
        sid = payload.session_id or f"pg-{agent_id}-{user['id']}"
        try:
            from emergentintegrations.llm.chat import UserMessage
            reply = await _llm(sys, sid).send_message(UserMessage(text=payload.message))
            reply = str(reply)
        except Exception as e:
            logger.error(f"agent chat failed: {e}")
            raise HTTPException(502, f"AI response failed: {e}")
        confidence = round(sources[0]["score"], 2) if sources else 0.0
        grounded = bool(sources) and confidence >= 0.12
        await db.agent_chats.insert_one({
            "id": str(uuid.uuid4()), "agent_id": agent_id, "question": payload.message,
            "confidence": confidence, "grounded": grounded,
            "created_at": datetime.now(timezone.utc).isoformat()})
        if not grounded:
            await _record_unknown(db, agent_id, payload.message)
        return {"reply": reply, "sources": [{"title": s["title"], "snippet": s["content"][:180], "score": s["score"]} for s in sources],
                "confidence": confidence, "grounded": grounded}

    # ---- Q&A training data ----
    @router.get("/{agent_id}/qa")
    async def list_qa(agent_id: str, user=Depends(get_current_user)):
        docs = await db.agent_docs.find({"agent_id": agent_id, "type": "qa"}).sort("created_at", -1).to_list(500)
        for d in docs:
            d.pop("_id", None)
        return docs

    @router.post("/{agent_id}/qa")
    async def add_qa(agent_id: str, payload: dict, user=Depends(get_current_user)):
        q = (payload.get("question") or "").strip()
        a = (payload.get("answer") or "").strip()
        if not q or not a:
            raise HTTPException(400, "Question and answer required")
        res = await _store_source(agent_id, f"Q&A: {q[:60]}", f"Q: {q}\nA: {a}", "qa")
        return {"ok": True, **res}

    # ---- Knowledge gaps / unknown questions ----
    @router.get("/{agent_id}/unknowns")
    async def list_unknowns(agent_id: str, user=Depends(get_current_user)):
        docs = await db.agent_unknowns.find({"agent_id": agent_id, "resolved": False}).sort("count", -1).to_list(500)
        for d in docs:
            d.pop("_id", None)
        return docs

    @router.post("/{agent_id}/unknowns/resolve")
    async def resolve_unknown(agent_id: str, payload: dict, user=Depends(get_current_user)):
        q = payload.get("question", "")
        answer = (payload.get("answer") or "").strip()
        if not answer:
            raise HTTPException(400, "Answer required to teach the agent")
        await _store_source(agent_id, f"Q&A: {q[:60]}", f"Q: {q}\nA: {answer}", "qa")
        await db.agent_unknowns.update_one({"agent_id": agent_id, "question": q.lower()}, {"$set": {"resolved": True}})
        return {"ok": True}

    @router.delete("/{agent_id}/unknowns")
    async def dismiss_unknown(agent_id: str, question: str, user=Depends(get_current_user)):
        await db.agent_unknowns.update_one({"agent_id": agent_id, "question": question.lower()}, {"$set": {"resolved": True}})
        return {"ok": True}

    # ---- Prompt versioning ----
    @router.get("/{agent_id}/prompt-versions")
    async def prompt_versions(agent_id: str, user=Depends(get_current_user)):
        docs = await db.agent_prompt_versions.find({"agent_id": agent_id}).sort("created_at", -1).to_list(50)
        for d in docs:
            d.pop("_id", None)
        return docs

    @router.post("/{agent_id}/prompt-versions/rollback")
    async def rollback_prompt(agent_id: str, payload: dict, user=Depends(get_current_user)):
        ver = await db.agent_prompt_versions.find_one({"id": payload.get("version_id"), "agent_id": agent_id})
        if not ver:
            raise HTTPException(404, "Version not found")
        agent = await db.ai_agents.find_one({"id": agent_id})
        await db.agent_prompt_versions.insert_one({
            "id": str(uuid.uuid4()), "agent_id": agent_id, "system_prompt": agent.get("system_prompt") or "",
            "created_at": datetime.now(timezone.utc).isoformat(), "created_by": user.get("name")})
        await db.ai_agents.update_one({"id": agent_id}, {"$set": {"system_prompt": ver["system_prompt"], "updated_at": datetime.now(timezone.utc).isoformat()}})
        return {"ok": True, "system_prompt": ver["system_prompt"]}

    # ---- Calls + QA review ----
    @router.get("/{agent_id}/calls")
    async def agent_calls(agent_id: str, user=Depends(get_current_user)):
        docs = await db.calls.find({"agent_id": agent_id}).sort("created_at", -1).to_list(200)
        for d in docs:
            d.pop("_id", None)
        return docs

    @router.post("/{agent_id}/calls/{call_id}/review")
    async def review_call(agent_id: str, call_id: str, payload: dict, user=Depends(get_current_user)):
        upd = {"qa_reviewed": True, "qa_reviewed_by": user.get("name")}
        if "rating" in payload:
            upd["qa_rating"] = int(payload["rating"])
        if "flagged" in payload:
            upd["qa_flagged"] = bool(payload["flagged"])
        if "note" in payload:
            upd["qa_note"] = payload["note"]
        r = await db.calls.find_one_and_update({"id": call_id, "agent_id": agent_id}, {"$set": upd}, return_document=True)
        if not r:
            raise HTTPException(404, "Call not found")
        return {"ok": True}

    # ---- Per-agent analytics ----
    @router.get("/{agent_id}/analytics")
    async def analytics(agent_id: str, user=Depends(get_current_user)):
        chats = await db.agent_chats.find({"agent_id": agent_id}).to_list(5000)
        calls = await db.calls.find({"agent_id": agent_id}).to_list(5000)
        unknowns = await db.agent_unknowns.count_documents({"agent_id": agent_id, "resolved": False})
        kdocs = await db.agent_docs.count_documents({"agent_id": agent_id})
        confs = [c.get("confidence", 0) for c in chats]
        outcomes = Counter(c.get("outcome", "unknown") for c in calls)
        top_q = Counter(c["question"].strip().lower() for c in chats if c.get("question"))
        grounded_rate = round(100 * sum(1 for c in chats if c.get("grounded")) / len(chats)) if chats else 0
        return {
            "chats": len(chats), "calls": len(calls), "knowledge_docs": kdocs,
            "unknown_questions": unknowns,
            "avg_confidence": round(sum(confs) / len(confs), 2) if confs else 0,
            "grounded_rate": grounded_rate,
            "outcomes": dict(outcomes),
            "top_questions": [{"q": q, "n": n} for q, n in top_q.most_common(8)],
        }

    return router


# ---------------- Grounded AI call generation (used by lead call engine) ----------------
async def generate_call(db, agent: dict, lead: dict, language: str) -> dict:
    """Generate a realistic, knowledge-grounded call transcript + summary using the agent's knowledge."""
    from emergentintegrations.llm.chat import UserMessage
    query = f"{lead.get('course') or ''} {lead.get('notes') or ''} fees demo enrollment"
    sources = await retrieve(db, agent["id"], query, k=6)
    context = "\n---\n".join(f"[{s['title']}] {s['content']}" for s in sources)
    sys = _agent_system_prompt(agent, context)
    instruction = (
        f"Simulate a short, realistic outbound phone call in {language}. "
        f"You are calling the lead named {lead.get('name','the customer')} "
        f"who is interested in '{lead.get('course') or 'our offering'}'. "
        "Have a natural back-and-forth (6-10 turns) where the customer asks 2-3 real questions "
        "answered strictly from your knowledge. Then return STRICT JSON only, no prose, with keys: "
        '{"transcript":[{"speaker":"AI"|"Customer","text":"..."}],'
        '"summary":"1-2 sentence summary","outcome":"interested|not_interested|callback|enrolled",'
        '"interest_score":0-100}.'
    )
    sid = f"call-{agent['id']}-{lead.get('id','x')}-{uuid.uuid4().hex[:6]}"
    raw = str(await _llm(sys, sid).send_message(UserMessage(text=instruction)))
    m = re.search(r"\{.*\}", raw, re.DOTALL)
    data = json.loads(m.group(0)) if m else {}
    return {
        "transcript": data.get("transcript") or [{"speaker": "AI", "text": raw[:500]}],
        "summary": data.get("summary") or "AI follow-up call completed.",
        "outcome": data.get("outcome") or "interested",
        "interest_score": data.get("interest_score"),
        "sources": [s["title"] for s in sources],
    }
